"""Attachment text extractors for the enrich stage (4B).

The core is ``extract_blob``: a pure function over one blob's bytes and metadata
that returns an ``AttachmentText`` result. It never raises for malformed input —
every failure maps to a status — and it never fetches remote resources. The
service runs it in an isolated subprocess; this module stays free of process or
I/O concerns so it can be tested directly.

Extractor selection is by ``detected_mime``, then the declared type, then the
filename extension. Legacy Office formats and nested ``message/rfc822`` are
``unsupported``.
"""

from __future__ import annotations

import io

from rosalind.domain.email import AttachmentStatus, AttachmentText
from rosalind.domain.email.text import clean_plain

_PDF_MIME = "application/pdf"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

_TEXT_MIMES = {"text/plain", "text/html", "text/csv", "text/markdown", "text/xml"}

_LEGACY_EXTENSIONS = {".doc", ".xls", ".ppt"}

# Extension fallback → extractor key, used only when no MIME resolved.
_EXTENSION_MAP = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pptx": "pptx",
    ".txt": "text",
    ".html": "text",
    ".htm": "text",
    ".csv": "text",
    ".md": "text",
}


def _pick_key(
    detected_mime: str | None, declared_mime: str | None, filename: str | None
) -> str | None:
    for mime in (detected_mime, declared_mime):
        if not mime:
            continue
        normalized = mime.split(";", 1)[0].strip().lower()
        if normalized == _PDF_MIME:
            return "pdf"
        if normalized == _DOCX_MIME:
            return "docx"
        if normalized == _XLSX_MIME:
            return "xlsx"
        if normalized == _PPTX_MIME:
            return "pptx"
        if normalized in _TEXT_MIMES or normalized.startswith("text/"):
            return "text"
    if filename:
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext in _LEGACY_EXTENSIONS:
            return "unsupported"
        if ext in _EXTENSION_MAP:
            return _EXTENSION_MAP[ext]
    return None


def extract_blob(
    *,
    blob_sha256: str,
    data: bytes,
    detected_mime: str | None,
    declared_mime: str | None,
    filename: str | None,
    max_output_chars: int,
) -> AttachmentText:
    key = _pick_key(detected_mime, declared_mime, filename)
    if key is None:
        return AttachmentText(
            blob_sha256=blob_sha256,
            status=AttachmentStatus.UNSUPPORTED,
            error="no extractable type",
        )
    if key == "unsupported":
        return AttachmentText(
            blob_sha256=blob_sha256,
            status=AttachmentStatus.UNSUPPORTED,
            error="legacy Office format not supported",
        )

    extractor = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "xlsx": _extract_xlsx,
        "pptx": _extract_pptx,
        "text": _extract_text,
    }[key]

    try:
        text, page_count = extractor(data)
    except _Encrypted as exc:
        return AttachmentText(
            blob_sha256=blob_sha256,
            status=AttachmentStatus.ENCRYPTED,
            error=str(exc),
        )
    except _Empty as exc:
        return AttachmentText(
            blob_sha256=blob_sha256,
            status=AttachmentStatus.NEEDS_OCR,
            error=str(exc),
        )
    except Exception as exc:  # noqa: BLE001 - untrusted input, never raise
        return AttachmentText(
            blob_sha256=blob_sha256,
            status=AttachmentStatus.FAILED,
            error=f"{type(exc).__name__}: {exc}",
        )

    text = clean_plain(text)
    truncated = False
    if len(text) > max_output_chars:
        text, truncated = _truncate(text, max_output_chars)

    if not text.strip():
        return AttachmentText(
            blob_sha256=blob_sha256,
            status=AttachmentStatus.EMPTY,
            method=key,
            page_count=page_count,
            truncated=truncated,
        )

    return AttachmentText(
        blob_sha256=blob_sha256,
        status=AttachmentStatus.DONE,
        text=text,
        method=key,
        page_count=page_count,
        truncated=truncated,
    )


def _truncate(text: str, cap: int) -> tuple[str, bool]:
    if len(text) <= cap:
        return text, False
    cut = text.rfind("\n", 0, cap)
    if cut <= 0:
        cut = cap
    return text[:cut], True


class _Encrypted(Exception):
    pass


class _Empty(Exception):
    pass


def _extract_text(data: bytes) -> tuple[str, int | None]:
    return data.decode("utf-8", errors="replace"), None


def _extract_pdf(data: bytes) -> tuple[str, int | None]:
    from pdfminer.high_level import extract_text
    from pdfminer.pdfdocument import PDFEncryptionError
    from pdfminer.pdfpage import PDFPage

    stream = io.BytesIO(data)
    try:
        page_count = sum(1 for _ in PDFPage.get_pages(stream))
    except PDFEncryptionError as exc:
        raise _Encrypted("PDF is encrypted") from exc
    except Exception:  # noqa: BLE001 - page count is best-effort
        page_count = None

    try:
        text = extract_text(io.BytesIO(data))
    except PDFEncryptionError as exc:
        raise _Encrypted("PDF is encrypted") from exc
    if not text.strip():
        raise _Empty("PDF has no text layer")
    return text, page_count


def _extract_docx(data: bytes) -> tuple[str, int | None]:
    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs]
    return "\n".join(paragraphs), None


def _extract_xlsx(data: bytes) -> tuple[str, int | None]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    # read_only + data_only: bounded memory, cached values instead of formulas.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    row_cap = 10_000
    cell_cap = 200
    try:
        for sheet in workbook.worksheets:
            parts.append(sheet.title)
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= row_cap:
                    break
                cells = ["" if v is None else str(v) for v in row[:cell_cap]]
                parts.append("\t".join(cells))
    finally:
        workbook.close()
    return "\n".join(parts), None


def _extract_pptx(data: bytes) -> tuple[str, int | None]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts), len(presentation.slides) or None
